# course_config_loader.py
# Author: Levester Williams
# 1 May 2026
#

from openpyxl import load_workbook
from typing import List, Tuple


def read_course_config_xlsx(path: str) -> List[Tuple[str, str, str]]:
    """
    Read course configuration from Excel.

    Parameters
    ----------
    path : str
        Path to XLSX file.

    Returns
    -------
    list of tuple
        List of (course_name, canvas_course_id, box_folder_id)
    """
    wb = load_workbook(path)
    ws = wb.active

    rows = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        course_name, course_id, box_folder_id = row

        if not course_id or not box_folder_id:
            print(f"Skipping row {i}: missing required fields")
            continue

        rows.append((
            str(course_name or "Unknown Course"),
            str(course_id),
            str(box_folder_id)
        ))

    return rows